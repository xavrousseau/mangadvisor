"""Bounded, offline XLSX -> CSV conversion for reading-library import only.

Never executes spreadsheet formulae and never uses an external parser service.
The existing ADD-ONLY CSV endpoint remains the single business import path.
"""
from __future__ import annotations

import base64
import binascii
import csv
import io

from fastapi import HTTPException
from openpyxl import load_workbook

MAX_XLSX_BYTES = 512_000
MAX_ROWS = 1_500
MAX_COLUMNS = 32
MAX_CSV_CHARS = 150_000


def xlsx_to_library_csv(content_b64: str) -> str:
    try:
        data = base64.b64decode(content_b64, validate=True)
    except (ValueError, binascii.Error) as exc:
        raise HTTPException(status_code=400, detail="Fichier Excel encodé invalide.") from exc
    if not data.startswith(b"PK\\x03\\x04") or len(data) > MAX_XLSX_BYTES:
        raise HTTPException(status_code=400, detail="Classeur XLSX invalide ou trop volumineux.")
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            if not workbook.worksheets:
                raise ValueError("Classeur Excel vide.")
            sheet = workbook.worksheets[0]
            if sheet.max_row and sheet.max_row > MAX_ROWS:
                raise ValueError("Trop de lignes dans le classeur.")
            if sheet.max_column and sheet.max_column > MAX_COLUMNS:
                raise ValueError("Trop de colonnes dans le classeur.")
            output = io.StringIO()
            writer = csv.writer(output, lineterminator="\\n")
            for index, row in enumerate(sheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS or len(row) > MAX_COLUMNS:
                    raise ValueError("Limite de taille Excel dépassée.")
                values = ["" if value is None else str(value) for value in row]
                if any(values):
                    writer.writerow(values)
                if output.tell() > MAX_CSV_CHARS:
                    raise ValueError("Contenu Excel trop volumineux.")
            return output.getvalue()
        finally:
            workbook.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=400, detail="Classeur Excel illisible ou hors limites.") from exc
